"""Unit tests for the grading, reference-mapping, and item-analysis logic.

These exercise the two trickiest correctness traps in this app: the pandas
"NA"-as-NaN gotcha (see data_loader.py), and the Exam-A<->Exam-B canonical
question/option mapping used for item analysis (see reference.py /
item_analysis.py).
"""
from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src.data_loader import load_grading_sheet, load_reference_sheet, WorkbookValidationError
from src.grading import build_correctness_mask, score_grading_sheet
from src.item_analysis import option_selection_breakdown, question_difficulty
from src.reference import build_answer_keys

EXAM_NUMBER = 1
SHIFT = 25  # self-inverse cyclic shift over 1..50, used to build a synthetic A<->B mapping


def _b_question(a_question: int) -> int:
    return ((a_question - 1 + SHIFT) % 50) + 1


def _reference_df() -> pd.DataFrame:
    """One row per (question, option) pair. Each question has two options:
    canonical (Exam A) letter "A" is always correct and reads "B" on Exam B;
    canonical letter "B" is always the distractor and reads "A" on Exam B.
    """
    rows = []
    for a_q in range(1, 51):
        b_q = _b_question(a_q)
        rows.append(
            {"A - Question": a_q, "A - Option": "A", "B - Question": b_q, "B - Option": "B", "Correct": 1}
        )
        rows.append(
            {"A - Question": a_q, "A - Option": "B", "B - Question": b_q, "B - Option": "A", "Correct": 0}
        )
    return pd.DataFrame(rows)


def _grading_df() -> pd.DataFrame:
    student_1 = {"Student": "Alice", "Exam A or B?": "A", **{q: "A" for q in range(1, 51)}}
    student_2 = {"Student": "Bob", "Exam A or B?": "B", **{q: "B" for q in range(1, 51)}}
    student_3 = {
        "Student": "Carol",
        "Exam A or B?": "A",
        **{q: "A" for q in range(1, 26)},
        **{q: "X" for q in range(26, 51)},
    }
    return pd.DataFrame([student_1, student_2, student_3])


def test_score_grading_sheet():
    keys = build_answer_keys(_reference_df())
    graded = score_grading_sheet(_grading_df(), keys)

    assert graded.set_index("Student")["Total"].to_dict() == {
        "Alice": 100,
        "Bob": 100,
        "Carol": 50,
    }


def test_build_answer_keys_mapping():
    keys = build_answer_keys(_reference_df())

    assert keys.key_a[1] == {"A"}
    assert keys.key_b[_b_question(1)] == {"B"}
    assert keys.is_correct("A", 1, "A")
    assert not keys.is_correct("A", 1, "B")
    assert keys.correct_letters("A", 1) == "A"
    assert keys.a_to_b[1] == _b_question(1)
    assert keys.b_to_a[_b_question(1)] == 1
    assert keys.canonical_id("B", _b_question(7)) == 7

    # A student's native letter, on either version, resolves to the same
    # canonical (Exam A) option identity.
    assert keys.canonical_option("A", 1, "A") == "A"
    assert keys.canonical_option("A", 1, "B") == "B"
    assert keys.canonical_option("B", _b_question(1), "B") == "A"
    assert keys.canonical_option("B", _b_question(1), "A") == "B"
    assert keys.canonical_option("A", 1, "Z") is None

    assert keys.b_option(1, "A") == "B"
    assert keys.canonical_options[1] == ["A", "B"]


def _reference_df_with_multi_correct_question(multi_a_q: int) -> pd.DataFrame:
    """Same as `_reference_df()`, except `multi_a_q` accepts both of its
    options as correct (both rows flagged `Correct = 1`) - simulating a
    thrown-out question with two accepted answers."""
    df = _reference_df()
    df.loc[(df["A - Question"] == multi_a_q) & (df["A - Option"] == "B"), "Correct"] = 1
    return df


def test_build_answer_keys_supports_multiple_correct_answers():
    keys = build_answer_keys(_reference_df_with_multi_correct_question(1))
    b_q1 = _b_question(1)

    assert keys.key_a[1] == {"A", "B"}
    assert keys.key_b[b_q1] == {"A", "B"}
    assert keys.is_correct("A", 1, "A")
    assert keys.is_correct("A", 1, "B")
    assert keys.is_correct("B", b_q1, "A")
    assert keys.is_correct("B", b_q1, "B")
    assert keys.correct_letters("A", 1) == "A / B"

    # Every other question is untouched - still single-answer.
    assert keys.key_a[2] == {"A"}


def test_multi_correct_answer_question_scores_either_option_as_correct():
    keys = build_answer_keys(_reference_df_with_multi_correct_question(1))

    grading_df = pd.DataFrame(
        [
            {"Student": "Pat", "Exam A or B?": "A", **{q: "A" for q in range(1, 51)}},
            # Frankie picks the *other* accepted option for Q1 only - should
            # still be scored correct there, unlike every other A-only question.
            {"Student": "Frankie", "Exam A or B?": "A", 1: "B", **{q: "A" for q in range(2, 51)}},
        ]
    )
    graded = score_grading_sheet(grading_df, keys)
    assert graded.set_index("Student")["Total"].to_dict() == {"Pat": 100, "Frankie": 100}

    mask = build_correctness_mask(grading_df, keys)
    frankie_index = grading_df.index[grading_df["Student"] == "Frankie"][0]
    assert mask.loc[frankie_index, 1]

    difficulty = question_difficulty(graded, keys)
    row_q1 = difficulty[difficulty["Exam A Question Number"] == 1].iloc[0]
    assert row_q1["Percent Correct"] == pytest.approx(100.0)
    assert row_q1["Exam A Answer"] == "A / B"

    breakdown = option_selection_breakdown(graded, keys)
    row_a = _breakdown_row(breakdown, 1, "A")
    row_b = _breakdown_row(breakdown, 1, "B")
    assert row_a["Is Correct"]
    assert row_b["Is Correct"]
    assert row_a["Count"] == 1
    assert row_b["Count"] == 1


def test_build_correctness_mask():
    grading_df = _grading_df()
    keys = build_answer_keys(_reference_df())
    mask = build_correctness_mask(grading_df, keys)

    carol_index = grading_df.index[grading_df["Student"] == "Carol"][0]
    assert mask.loc[carol_index, list(range(1, 26))].all()
    assert not mask.loc[carol_index, list(range(26, 51))].any()


def test_question_difficulty_combines_versions_and_carries_tooltip_fields():
    keys = build_answer_keys(_reference_df())
    graded = score_grading_sheet(_grading_df(), keys)
    difficulty = question_difficulty(graded, keys)

    # sorted ascending by Exam A Question Number, not by Percent Correct
    assert difficulty["Exam A Question Number"].tolist() == list(range(1, 51))

    by_question = difficulty.set_index("Exam A Question Number")["Percent Correct"].to_dict()
    easy_questions = {q for q in range(1, 26)}
    hard_questions = {q for q in range(26, 51)}

    assert all(by_question[q] == pytest.approx(100.0) for q in easy_questions)
    assert all(by_question[q] == pytest.approx(200 / 3) for q in hard_questions)

    row_q1 = difficulty[difficulty["Exam A Question Number"] == 1].iloc[0]
    assert row_q1["Exam A Answer"] == "A"
    assert row_q1["Exam B Question Number"] == _b_question(1)
    assert row_q1["Exam B Answer"] == "B"


def _breakdown_row(breakdown: pd.DataFrame, a_q: int, option: str) -> pd.Series:
    match = breakdown[
        (breakdown["Exam A Question Number"] == a_q) & (breakdown["Option"] == option)
    ]
    return match.iloc[0]


def test_option_selection_breakdown_combines_versions_and_flags_correct_option():
    keys = build_answer_keys(_reference_df())
    graded = score_grading_sheet(_grading_df(), keys)
    breakdown = option_selection_breakdown(graded, keys)

    # Every canonical question carries all option buckets that exist for it
    # (A/B plus NA here, since the synthetic reference only defines 2 options
    # per question), even ones nobody picked, so stacked-bar segments align.
    assert len(breakdown) == 50 * 3

    # Question 1 (easy): Alice (A) and Carol (A) both picked native "A";
    # Bob (B) picked native "B" — all three map to canonical option "A",
    # which is correct. Combined across both versions in one row.
    row = _breakdown_row(breakdown, 1, "A")
    assert row["Count"] == 3
    assert row["A Count"] == 2
    assert row["B Count"] == 1
    assert row["Total"] == 3
    assert row["Percent"] == pytest.approx(100.0)
    assert row["Is Correct"]
    assert row["B Option"] == "B"

    row = _breakdown_row(breakdown, 1, "B")
    assert row["Count"] == 0

    # Question 26 (hard): Alice and Bob are correct (canonical "A"); Carol's
    # stray "X" answer lands in "NA" (graded incorrect) rather than dropped.
    correct_row = _breakdown_row(breakdown, 26, "A")
    assert correct_row["Count"] == 2
    assert correct_row["Total"] == 3
    assert correct_row["Is Correct"]

    na_row = _breakdown_row(breakdown, 26, "NA")
    assert na_row["Count"] == 1
    assert na_row["A Count"] == 1
    assert na_row["B Count"] == 0
    assert na_row["Total"] == 3
    assert na_row["Percent"] == pytest.approx(100 / 3)
    assert not na_row["Is Correct"]


def test_load_grading_sheet_preserves_literal_na():
    """Blank answers are pre-filled by the teacher as the literal string "NA".

    pandas treats "NA" as a default missing-value token on read_excel, so
    without keep_default_na=False this would silently become NaN instead of
    the literal string the scoring logic compares against.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Exam {EXAM_NUMBER} Grading"

    header = ["Student", "Total", "Exam A or B?"] + list(range(1, 51))
    ws.append(header)

    row = ["Dana", 0, "A"] + ["A"] * 50
    row[header.index(1)] = "NA"  # blank answer to question 1
    ws.append(row)

    with tempfile.TemporaryDirectory() as tmp_dir:
        path = Path(tmp_dir) / "workbook.xlsx"
        wb.save(path)

        df = load_grading_sheet(str(path), EXAM_NUMBER)

    assert df.loc[0, 1] == "NA"
    assert "Total" not in df.columns


def test_load_grading_sheet_works_without_total_column():
    """Total is optional in the source workbook, not required."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Exam {EXAM_NUMBER} Grading"

    header = ["Student", "Exam A or B?"] + list(range(1, 51))  # no Total column
    ws.append(header)
    ws.append(["Erin", "A"] + ["A"] * 50)

    with tempfile.TemporaryDirectory() as tmp_dir:
        path = Path(tmp_dir) / "workbook.xlsx"
        wb.save(path)

        grading_df = load_grading_sheet(str(path), EXAM_NUMBER)

    assert "Total" not in grading_df.columns

    keys = build_answer_keys(_reference_df())
    graded = score_grading_sheet(grading_df, keys)
    assert graded.set_index("Student")["Total"].to_dict() == {"Erin": 100}


def test_load_grading_sheet_missing_sheet_raises():
    wb = openpyxl.Workbook()
    wb.active.title = "Some Other Sheet"

    with tempfile.TemporaryDirectory() as tmp_dir:
        path = Path(tmp_dir) / "workbook.xlsx"
        wb.save(path)

        with pytest.raises(WorkbookValidationError):
            load_grading_sheet(str(path), EXAM_NUMBER)


def _write_reference_sheet(rows: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Exam {EXAM_NUMBER} Reference"
    ws.append(["A - Question", "A - Option", "B - Question", "B - Option", "Correct"])
    for row in rows:
        ws.append([row["A - Question"], row["A - Option"], row["B - Question"], row["B - Option"], row["Correct"]])

    tmp_dir = tempfile.mkdtemp()
    path = Path(tmp_dir) / "workbook.xlsx"
    wb.save(path)
    return path


def test_load_reference_sheet_wrong_question_count_raises():
    # Only 1 distinct question, should be 50.
    path = _write_reference_sheet(
        [{"A - Question": 1, "A - Option": "A", "B - Question": 1, "B - Option": "A", "Correct": 1}]
    )

    with pytest.raises(WorkbookValidationError):
        load_reference_sheet(str(path), EXAM_NUMBER)


def test_load_reference_sheet_bad_correct_flag_count_raises():
    rows = []
    for a_q in range(1, 51):
        # Question 1 has no correct option flagged; every other question has one.
        correct = 0 if a_q == 1 else 1
        rows.append({"A - Question": a_q, "A - Option": "A", "B - Question": a_q, "B - Option": "A", "Correct": correct})
    path = _write_reference_sheet(rows)

    with pytest.raises(WorkbookValidationError):
        load_reference_sheet(str(path), EXAM_NUMBER)


def test_load_reference_sheet_accepts_text_columns_and_preserves_none_text():
    """Question Text / Answer Text ride along with the mapping columns, and
    literal text like "None" must not be turned into NaN by pandas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Exam {EXAM_NUMBER} Reference"
    ws.append(
        ["A - Question", "A - Option", "B - Question", "B - Option", "Correct", "Question Text", "Answer Text"]
    )
    for row in _reference_df().itertuples(index=False):
        answer_text = "None" if (row[0] == 1 and row[1] == "B") else f"Answer {row[1]}"
        ws.append([*row, f"Question {row[0]}", answer_text])

    with tempfile.TemporaryDirectory() as tmp_dir:
        path = Path(tmp_dir) / "workbook.xlsx"
        wb.save(path)

        df = load_reference_sheet(str(path), EXAM_NUMBER)

    assert {"Question Text", "Answer Text"} <= set(df.columns)
    none_row = df[(df["A - Question"] == 1) & (df["A - Option"] == "B")].iloc[0]
    assert none_row["Answer Text"] == "None"
    build_answer_keys(df)
